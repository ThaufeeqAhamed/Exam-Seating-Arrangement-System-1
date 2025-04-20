from django.shortcuts import render, get_object_or_404
from rest_framework import viewsets, status, permissions
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated, AllowAny, IsAdminUser
from rest_framework_simplejwt.views import TokenObtainPairView
from django.contrib.auth import get_user_model
import pandas as pd
from io import BytesIO
from openpyxl import load_workbook
from rest_framework.exceptions import ValidationError
from rest_framework.parsers import MultiPartParser, FormParser
from django.db.models import Q
from django.conf import settings
import os

from .models import Room, Student, Exam, SeatingArrangement, Timetable
from .utils import process_pdf_timetable
from .serializers import (
    UserSerializer, UserCreateSerializer, RoomSerializer,
    StudentSerializer, ExamSerializer, SeatingArrangementSerializer,
    ExamWithSeatingSerializer, TimetableSerializer
)
from .utils import (
    generate_seating_arrangement,
)

User = get_user_model()

class CustomTokenObtainPairView(TokenObtainPairView):
    def validate_input(self, data):
        errors = {}
        if not data.get('email'):
            errors['email'] = 'Email is required'
        if not data.get('password'):
            errors['password'] = 'Password is required'
        return errors

    def post(self, request, *args, **kwargs):
        errors = self.validate_input(request.data)
        if errors:
            return Response(
                {'errors': errors},
                status=status.HTTP_400_BAD_REQUEST
            )

        try:
            try:
                email = request.data.get('email')
                if not email:
                    raise ValidationError('Email is required')
                user = User.objects.get(email=email)
            except User.DoesNotExist:
                return Response(
                    {'error': 'Invalid credentials'},
                    status=status.HTTP_401_UNAUTHORIZED
                )

            response = super().post(request, *args, **kwargs)
            if response.status_code == 200:
                response.data.update({
                    'user': UserSerializer(user).data,
                    'message': 'Login successful'
                })
            return response

        except Exception as e:
            return Response(
                {'error': 'Authentication failed', 'detail': str(e)},
                status=status.HTTP_401_UNAUTHORIZED
            )

class UserViewSet(viewsets.ModelViewSet):
    queryset = User.objects.all()
    serializer_class = UserSerializer
    permission_classes = [IsAuthenticated]

    def get_serializer_class(self):
        if self.action == 'create':
            return UserCreateSerializer
        return UserSerializer

    def get_permissions(self):
        if self.action == 'create':
            return [AllowAny()]
        return super().get_permissions()

    @action(detail=False, methods=['get'])
    def me(self, request):
        serializer = self.get_serializer(request.user)
        return Response(serializer.data)

class RoomViewSet(viewsets.ModelViewSet):
    queryset = Room.objects.all()
    serializer_class = RoomSerializer
    permission_classes = [IsAuthenticated]
    filterset_fields = ['floor_number', 'capacity']
    search_fields = ['room_id']
    ordering_fields = ['room_id', 'floor_number', 'capacity']
    ordering = ['floor_number', 'room_id']

    def perform_create(self, serializer):
        if serializer.validated_data['capacity'] <= 0:
            raise ValidationError({'capacity': 'Capacity must be greater than 0'})
        if serializer.validated_data['number_of_benches'] <= 0:
            raise ValidationError({'number_of_benches': 'Number of benches must be greater than 0'})
        serializer.save()

    def perform_update(self, serializer):
        if serializer.validated_data.get('capacity', 1) <= 0:
            raise ValidationError({'capacity': 'Capacity must be greater than 0'})
        if serializer.validated_data.get('number_of_benches', 1) <= 0:
            raise ValidationError({'number_of_benches': 'Number of benches must be greater than 0'})
        serializer.save()

class StudentViewSet(viewsets.ModelViewSet):
    queryset = Student.objects.all().select_related()
    serializer_class = StudentSerializer
    permission_classes = [IsAuthenticated]
    filterset_fields = ['department', 'year_of_study', 'semester']
    search_fields = ['roll_number', 'name']
    ordering_fields = ['roll_number', 'name', 'year_of_study']
    ordering = ['roll_number']

    def get_permissions(self):
        if self.action == 'clear':
            return [IsAuthenticated()]
        return super().get_permissions()

    def validate_student_data(self, data, instance=None):
        errors = {}
        required_fields = ['roll_number', 'name', 'department']

        for field in required_fields:
            value = data.get(field)
            if not value and (not instance or not getattr(instance, field)):
                errors[field] = f'{field.title().replace("_", " ")} is required'

        return errors

    def perform_create(self, serializer):
        errors = self.validate_student_data(serializer.validated_data)
        if errors:
            raise ValidationError(errors)
        serializer.save()

    def perform_update(self, serializer):
        errors = self.validate_student_data(serializer.validated_data, self.get_object())
        if errors:
            raise ValidationError(errors)
        serializer.save()

    @action(detail=False, methods=['delete'])
    def clear(self, request):
        try:
            count = Student.objects.count()
            Student.objects.all().delete()
            return Response({
                'message': f'Successfully deleted {count} students',
                'count': count
            }, status=status.HTTP_200_OK)
        except Exception as e:
            return Response({
                'error': f'Failed to clear students: {str(e)}'
            }, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

    @action(detail=False, methods=['post'])
    def upload_excel(self, request):
        if 'file' not in request.FILES:
            return Response({'error': 'No file provided'}, status=status.HTTP_400_BAD_REQUEST)

        file = request.FILES['file']
        try:
            wb = load_workbook(file)
            sheet = wb.active
            required_columns = {'roll_number', 'name', 'department', 'year_of_study', 'semester',
                              'subject_code', 'seat_number'}
            headers = {cell.value for cell in sheet[1]}
            if not required_columns.issubset(headers):
                raise ValidationError(f"Missing required columns: {required_columns - headers}")

            created, updated, errors = 0, 0, []

            for row in sheet.iter_rows(min_row=2, values_only=True):
                try:
                    if not row[0]:
                        continue

                    student, created_flag = Student.objects.update_or_create(
                        roll_number=row[0],
                        defaults={
                            'name': row[1],
                            'department': row[2],
                            'year_of_study': row[3],
                            'semester': row[4],
                            'subject_code': row[5],
                            'seat_number': row[6],
                        }
                    )
                    if created_flag:
                        created += 1
                    else:
                        updated += 1
                except Exception as e:
                    errors.append(f"Row {row} - Error: {str(e)}")

            if created == 0 and updated == 0:
                return Response({
                    "error": "No students were processed. Please check your file format.",
                    "details": errors
                }, status=status.HTTP_400_BAD_REQUEST)

            return Response({
                "message": "Student data uploaded successfully.",
                "created": created,
                "updated": updated,
                "errors": errors
            }, status=status.HTTP_200_OK)
        except Exception as e:
            return Response({'error': str(e)}, status=status.HTTP_400_BAD_REQUEST)

class ExamViewSet(viewsets.ModelViewSet):
    queryset = Exam.objects.all().select_related()
    serializer_class = ExamSerializer
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]

    @action(detail=False, methods=['post'])
    def upload(self, request):
        if 'pdf_file' not in request.FILES:
            return Response({'error': 'No PDF file provided'}, status=status.HTTP_400_BAD_REQUEST)

        pdf_file = request.FILES['pdf_file']
        exam_type = request.data.get('exam_type', 'semester')

        try:
            # Process the PDF file and extract timetable entries
            entries, errors = process_pdf_timetable(pdf_file)

            if not entries:
                return Response({
                    'error': 'Failed to extract timetable data from PDF',
                    'details': errors
                }, status=status.HTTP_400_BAD_REQUEST)

            # Add exam type to each entry
            for entry in entries:
                entry['exam_type'] = exam_type

            response_data = {
                'message': 'Timetable processed successfully',
                'entries': entries
            }

            # Include any non-critical errors as warnings
            if errors:
                response_data['warnings'] = errors

            return Response(response_data, status=status.HTTP_200_OK)

        except Exception as e:
            return Response(
                {'error': f'Failed to process timetable: {str(e)}'},
                status=status.HTTP_400_BAD_REQUEST
            )

    def get_queryset(self):
        queryset = super().get_queryset()
        if self.action in ['retrieve', 'seating_details']:
            return queryset.prefetch_related('seating_arrangements__student')
        return queryset

    def get_serializer_class(self):
        if self.action in ['retrieve', 'seating_details']:
            return ExamWithSeatingSerializer
        return ExamSerializer

    def validate_exam_data(self, data):
        errors = {}
        required_fields = ['name', 'date', 'time', 'room']
        for field in required_fields:
            if not data.get(field):
                errors[field] = f'{field.title()} is required'
        return errors

    def perform_create(self, serializer):
        errors = self.validate_exam_data(serializer.validated_data)
        if errors:
            raise ValidationError(errors)
        serializer.save()

    @action(detail=True, methods=['post'])
    def generate_seating(self, request, pk=None):
        exam = self.get_object()

        if SeatingArrangement.objects.filter(exam=exam).exists():
            return Response(
                {'error': 'Seating arrangement already exists for this exam'},
                status=status.HTTP_400_BAD_REQUEST
            )

        if not exam.room:
            return Response(
                {'error': 'No room assigned to this exam'},
                status=status.HTTP_400_BAD_REQUEST
            )

        try:
            result = generate_seating_arrangement(pk)
            if 'error' in result:
                return Response({'error': result['error']}, status=status.HTTP_400_BAD_REQUEST)
            return Response({'message': result['success']}, status=status.HTTP_200_OK)
        except Exception as e:
            return Response(
                {'error': f'Failed to generate seating arrangement: {str(e)}'},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )

    @action(detail=True, methods=['get'])
    def seating_details(self, request, pk=None):
        exam = self.get_object()
        serializer = ExamWithSeatingSerializer(exam)
        return Response(serializer.data)

    @action(detail=True, methods=['get'])
    def department_seating(self, request, pk=None):
        exam = self.get_object()
        department = request.user.department

        if not department:
            return Response(
                {'error': 'You must have a department assigned to view department-specific seating'},
                status=status.HTTP_400_BAD_REQUEST
            )

        arrangements = SeatingArrangement.objects.filter(
            exam=exam,
            student__department=department
        )

        serializer = SeatingArrangementSerializer(arrangements, many=True)
        return Response({
            'exam': ExamSerializer(exam).data,
            'seating_arrangements': serializer.data
        })

class TimetableViewSet(viewsets.ModelViewSet):
    queryset = Timetable.objects.all()
    serializer_class = TimetableSerializer
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]
    filterset_fields = ['exam_type', 'status', 'department']
    search_fields = ['subject_code', 'subject_name']
    ordering_fields = ['date', 'start_time']
    ordering = ['date', 'start_time']

    @action(detail=False, methods=['post'], url_path="upload_file")
    def upload_file(self, request):
        if 'file' not in request.FILES:
            print("No file key in request.FILES")
            return Response({
                'error': 'No file provided',
                'details': 'Please select a PDF file to upload'
            }, status=status.HTTP_400_BAD_REQUEST)

        file = request.FILES['file']

        if not file.name.endswith('.pdf'):
            return Response({
                'error': 'Invalid file format',
                'details': 'Please upload a PDF file only. Other file formats are not supported.'
            }, status=status.HTTP_400_BAD_REQUEST)

        try:
            # Process the PDF file and extract timetable entries
            entries, errors = process_pdf_timetable(file)

            if not entries:
                error_message = 'Failed to extract timetable data from PDF'
                if errors:
                    error_message += f": {errors[0]}"
                return Response({
                    'error': error_message,
                    'details': errors
                }, status=status.HTTP_400_BAD_REQUEST)

            created, updated = 0, 0
            for entry in entries:
                try:
                    timetable, created_flag = Timetable.objects.update_or_create(
                        subject_code=entry['subject_code'],
                        exam_date=entry['exam_date'],
                        start_time=entry['exam_time'],
                        defaults={
                            'subject_name': entry['subject_name'],
                            'exam_type': entry.get('exam_type', 'semester'),
                            'department': entry.get('department', ''),
                            'status': 'scheduled'
                        }
                    )

                    if created_flag:
                        created += 1
                    else:
                        updated += 1

                except Exception as e:
                    errors.append(f"Entry {entry} - Error: {str(e)}")

            response_data = {
                "message": "Timetable data uploaded successfully.",
                "created": created,
                "updated": updated
            }

            if errors:
                response_data["warnings"] = errors

            return Response(response_data, status=status.HTTP_200_OK)

        except Exception as e:
            return Response({'error': str(e)}, status=status.HTTP_400_BAD_REQUEST)

class SeatingArrangementViewSet(viewsets.ModelViewSet):
    queryset = SeatingArrangement.objects.select_related('exam', 'student').all()
    serializer_class = SeatingArrangementSerializer
    permission_classes = [IsAuthenticated]
    filterset_fields = ['exam', 'bench_number', 'position']
    search_fields = ['student__name', 'student__roll_number']
    ordering_fields = ['bench_number', 'position', 'student__roll_number']
    ordering = ['bench_number', 'position']

    def get_queryset(self):
        user = self.request.user
        base_queryset = SeatingArrangement.objects.select_related('exam', 'student')

        if user.user_type == 'admin':
            return base_queryset.all()

        elif user.user_type == 'faculty' and user.department:
            return base_queryset.filter(student__department=user.department)

        return SeatingArrangement.objects.none()

    def perform_create(self, serializer):
        exam = serializer.validated_data.get('exam')
        student = serializer.validated_data.get('student')
        bench_number = serializer.validated_data.get('bench_number')
        position = serializer.validated_data.get('position')

        if SeatingArrangement.objects.filter(exam=exam, student=student).exists():
            raise ValidationError({'student': 'Student already has a seating arrangement for this exam'})

        if SeatingArrangement.objects.filter(exam=exam, bench_number=bench_number, position=position).exists():
            raise ValidationError({'position': 'This seat is already occupied'})

        if bench_number > exam.room.number_of_benches:
            raise ValidationError({'bench_number': 'Invalid bench number for this room'})

        serializer.save()
